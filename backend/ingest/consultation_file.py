import csv
from io import TextIOWrapper
from pathlib import Path
from typing import BinaryIO

from django.db import transaction
from openpyxl import load_workbook

from consultations.models import Consultation, Question, Respondent, Response


def read_consultation_file(
    file: BinaryIO, filename: str
) -> tuple[list[str], list[list[str | None]]]:
    extension = Path(filename).suffix.lower()
    file.seek(0)

    if extension == ".csv":
        wrapper = TextIOWrapper(file, encoding="utf-8-sig", newline="")
        try:
            rows = list(csv.reader(wrapper))
        finally:
            wrapper.detach()
    else:
        workbook = load_workbook(file, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet is None:
                raise ValueError("XLSX file has no visible worksheet.")
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()

    file.seek(0)
    if not rows:
        raise ValueError("File must contain a header row.")

    questions = ["" if value is None else str(value) for value in rows[0]]
    if not questions:
        raise ValueError("File must contain at least one question column.")

    response_rows = [
        [None if value is None or value == "" else str(value) for value in row[: len(questions)]]
        + [None] * max(0, len(questions) - len(row))
        for row in rows[1:]
    ]
    return questions, response_rows


@transaction.atomic
def import_consultation_data(
    consultation: Consultation, question_texts: list[str], response_rows: list[list[str | None]]
) -> tuple[int, int]:
    next_question_number = (
        consultation.question_set.order_by("-number").values_list("number", flat=True).first() or 0
    )
    questions = Question.objects.bulk_create(
        [
            Question(
                consultation=consultation,
                text=text,
                number=next_question_number + index,
                type=None,
            )
            for index, text in enumerate(question_texts, start=1)
        ]
    )
    respondents = Respondent.objects.bulk_create(
        [
            Respondent(consultation=consultation, themefinder_id=index)
            for index, _ in enumerate(response_rows, start=1)
        ]
    )
    responses = Response.objects.bulk_create(
        [
            Response(question=question, respondent=respondent, text=text)
            for respondent, row in zip(respondents, response_rows, strict=True)
            for question, text in zip(questions, row, strict=True)
        ],
        batch_size=1000,
    )
    return len(questions), len(responses)
