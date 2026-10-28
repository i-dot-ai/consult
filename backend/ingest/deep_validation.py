import csv
import re
from collections.abc import Sequence
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from consultations.models import Consultation

ILLEGAL_CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
ILLEGAL_CONTROL_BYTES = re.compile(rb"[\x00-\x08\x0b\x0c\x0e-\x1f]")
ILLEGAL_CONTROL_MARKER = "__CONSULT_ILLEGAL_CONTROL_CHARACTER__"
XLSX_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
BINARY_SIGNATURES = (
    b"%PDF-",
    b"\x89PNG\r\n\x1a\n",
    b"GIF87a",
    b"GIF89a",
    b"\xff\xd8\xff",
    b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",
    *XLSX_SIGNATURES,
)


def _citizen_space_rows(rows: Sequence[Sequence[object]]) -> tuple[int, list[int]]:
    return 1, list(range(2, len(rows) + 1))


def _qualtrics_rows(rows: Sequence[Sequence[object]]) -> tuple[int, list[int]]:
    return 2, list(range(4, len(rows) + 1))


ROW_SELECTORS = {
    str(Consultation.DataSource.CITIZEN_SPACE): _citizen_space_rows,
    str(Consultation.DataSource.QUALTRICS): _qualtrics_rows,
}


def _error(code: str, message: str, cell: str | None = None) -> dict:
    result = {"code": code, "message": message}
    if cell is not None:
        result["cell"] = cell
    return result


def _result(rejection_type: str, error: dict) -> dict:
    return {
        "valid": False,
        "rejection_type": rejection_type,
        "message": f"File failed deep validation with a {rejection_type} rejection.",
        "errors": [error],
    }


def _hard_rejection(code: str, message: str, cell: str | None = None) -> dict:
    return _result("hard", _error(code, message, cell))


def _read_and_validate_csv(content: bytes) -> list[list[str]] | None:
    if content.startswith(BINARY_SIGNATURES):
        return None

    rows = list(csv.reader(StringIO(content.decode("utf-8-sig"), newline="")))
    populated_row_lengths = {len(row) for row in rows if row}
    if len(populated_row_lengths) > 1:
        return None
    return rows


def _is_xlsx(content: bytes) -> bool:
    if not content.startswith(XLSX_SIGNATURES):
        return False
    try:
        with ZipFile(BytesIO(content)) as archive:
            names = set(archive.namelist())
    except BadZipFile:
        return False
    return "[Content_Types].xml" in names and "xl/workbook.xml" in names


def _read_xlsx(content: bytes) -> list[list[object]]:
    sanitised_content = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(sanitised_content, "w") as destination:
        for entry in source.infolist():
            entry_content = source.read(entry.filename)
            if entry.filename.endswith(".xml"):
                entry_content = ILLEGAL_CONTROL_BYTES.sub(
                    ILLEGAL_CONTROL_MARKER.encode(), entry_content
                )
            destination.writestr(entry, entry_content)

    workbook = load_workbook(BytesIO(sanitised_content.getvalue()), read_only=True, data_only=False)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("XLSX file has no visible worksheet.")
        return [list(row) for row in sheet.iter_rows(values_only=False)]
    finally:
        workbook.close()


def _cell_value(cell: object) -> object:
    return getattr(cell, "value", cell)


def _cell_reference(cell: object, row_number: int, column_number: int) -> str:
    return getattr(cell, "coordinate", f"{get_column_letter(column_number)}{row_number}")


def _is_formula(cell: object) -> bool:
    value = _cell_value(cell)
    return getattr(cell, "data_type", None) == "f" or (
        isinstance(value, str) and value.startswith("=")
    )


def _is_empty(row: Sequence[object]) -> bool:
    return not row or all(_cell_value(cell) in (None, "") for cell in row)


def _selected_rows(
    rows: Sequence[Sequence[object]], data_source: str
) -> tuple[list[tuple[int, Sequence[object]]], dict | None]:
    selector = ROW_SELECTORS.get(data_source)
    if selector is None:
        return [], _hard_rejection(
            "unsupported_data_source",
            "Consultation data source must be Citizen Space or Qualtrics.",
        )

    question_row_number, response_row_numbers = selector(rows)
    if len(rows) < question_row_number:
        return [], _hard_rejection("empty_questions", "The questions row is missing or empty.")
    if not response_row_numbers or len(rows) < response_row_numbers[0]:
        return [], _hard_rejection("empty_responses", "The file does not contain a response row.")

    selected = [(question_row_number, rows[question_row_number - 1])]
    selected.extend(
        (row_number, rows[row_number - 1])
        for row_number in response_row_numbers
        if row_number <= len(rows)
    )
    return selected, None


def _validate_rows(rows: Sequence[Sequence[object]], data_source: str) -> dict:
    selected_rows, selection_error = _selected_rows(rows, data_source)
    if selection_error:
        return selection_error

    soft_errors = []
    question_count = 0
    response_count = 0
    blank_cell_count = 0
    for selected_index, (row_number, row) in enumerate(selected_rows):
        row_type = "questions" if selected_index == 0 else "responses"
        if _is_empty(row):
            return _hard_rejection(
                f"empty_{row_type}",
                f"The {row_type} row is empty.",
            )

        if selected_index > 0:
            response_count += 1

        for column_number, cell in enumerate(row, start=1):
            value = _cell_value(cell)
            reference = _cell_reference(cell, row_number, column_number)
            if value in (None, ""):
                blank_cell_count += 1
            elif selected_index == 0:
                question_count += 1
            if isinstance(value, str) and (
                ILLEGAL_CONTROL_CHARACTERS.search(value) or ILLEGAL_CONTROL_MARKER in value
            ):
                return _hard_rejection(
                    "illegal_control_character",
                    "The cell contains an unreadable character.",
                    reference,
                )
            if _is_formula(cell):
                soft_errors.append(
                    _error(
                        "formula",
                        "Formulas are not allowed in question or response cells.",
                        reference,
                    )
                )

    if soft_errors:
        return {
            "valid": False,
            "rejection_type": "soft",
            "message": "File failed deep validation with a list of following soft rejection errors.",
            "errors": soft_errors,
            "question_count": question_count,
            "response_count": response_count,
            "blank_cell_count": blank_cell_count,
        }
    return {
        "valid": True,
        "rejection_type": None,
        "message": "File passed deep validation.",
        "errors": [],
        "question_count": question_count,
        "response_count": response_count,
        "blank_cell_count": blank_cell_count,
    }


def deep_validate_file(content: bytes, filename: str, data_source: str) -> dict:
    extension = Path(filename).suffix.lower()
    rows: Sequence[Sequence[object]]
    try:
        if extension == ".xlsx":
            if not _is_xlsx(content):
                return _hard_rejection(
                    "invalid_file_type", "File content is not a valid XLSX workbook."
                )
            rows = _read_xlsx(content)
        elif extension == ".csv":
            csv_rows = _read_and_validate_csv(content)
            if csv_rows is None:
                return _hard_rejection(
                    "invalid_file_type", "File content is not a valid CSV text file."
                )
            rows = csv_rows
        else:
            return _hard_rejection("invalid_file_type", "File must be a CSV or XLSX file.")
    except (BadZipFile, InvalidFileException, UnicodeDecodeError, ValueError, csv.Error) as error:
        return _hard_rejection("unreadable_file", f"File could not be read: {error}")

    return _validate_rows(rows, data_source)
